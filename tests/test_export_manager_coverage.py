"""Contratti sugli output reali che hanno sostituito il vecchio ExportManager."""

import json
import xml.etree.ElementTree as ET

from docx import Document
from PIL import Image
from PySide6.QtCore import QCoreApplication
from pypdf import PdfReader

from src.gedcom_factory import GedcomGenerator
from src.image_saver import save_image_variants
from src.ocr_processor import AdvancedOCRWorker
from src.pdf_utils import create_pdf_from_images
from src.translation_processor import TranslationWorker
import src.translation_processor as translation_processor
import src.translation_dialog as translation_dialog


class _TextBox:
    def __init__(self, text):
        self._text = text

    def toPlainText(self):
        return self._text


class _TranslationDialogStub:
    def __init__(self, text):
        self.txt_dest = _TextBox(text)

    def gm(self, text):
        return text


def test_current_output_pipeline_image_pdf_ocr_translation(tmp_path, monkeypatch):
    """Attraversa i writer effettivi senza dipendere da servizi di rete."""
    QCoreApplication.instance() or QCoreApplication([])
    source = tmp_path / "registro_sintetico.png"
    image = Image.new("RGB", (640, 480), "white")
    image.save(source)

    metadata = {
        "Title": "Registro sintetico",
        "_json": json.dumps({"source": "pre-rc", "canvas": 1}),
    }
    save_image_variants(
        image,
        str(tmp_path),
        "registro_canvas_1",
        ["PNG", "JPEG", "TIFF"],
        meta=metadata,
    )
    variants = [
        tmp_path / "registro_canvas_1.png",
        tmp_path / "registro_canvas_1.jpg",
        tmp_path / "registro_canvas_1.tif",
    ]
    for path in variants:
        with Image.open(path) as saved:
            saved.load()
            assert saved.size == (640, 480)
    assert json.loads(
        (tmp_path / "registro_canvas_1.json").read_text(encoding="utf-8")
    ) == {"source": "pre-rc", "canvas": 1}

    pdf_path = tmp_path / "registro.pdf"
    assert create_pdf_from_images([str(variants[0])], str(pdf_path)) == str(pdf_path)
    assert len(PdfReader(str(pdf_path)).pages) == 1
    assert not (tmp_path / "registro.pdf.progress.json").exists()

    transcription = (
        "ATTO DI NASCITA\nGiovanni Rossi, Trento, 4 marzo 1882.\n"
        "Padre Luigi Rossi. Madre Maria Bianchi."
    )
    ocr = AdvancedOCRWorker(
        "OpenAI",
        "test-key",
        ["txt", "docx", "xml"],
        str(tmp_path),
    )
    ocr.api_keys = ["test-key"]
    monkeypatch.setattr(ocr, "_transcribe_image", lambda *_: transcription)
    ocr.process_file(str(source))

    ocr_txt = tmp_path / "registro_sintetico_trascrizione.txt"
    ocr_docx = tmp_path / "registro_sintetico_trascrizione.docx"
    ocr_xml = tmp_path / "registro_sintetico_trascrizione.xml"
    assert ocr_txt.read_text(encoding="utf-8") == transcription
    assert transcription in "\n".join(p.text for p in Document(ocr_docx).paragraphs)
    assert ET.parse(ocr_xml).getroot().tag.endswith("TEI")

    translated = (
        "BIRTH RECORD\nGiovanni Rossi, Trento, 4 March 1882.\n"
        "Father Luigi Rossi. Mother Maria Bianchi."
    )
    captured = []
    worker = TranslationWorker("OpenAI", "test-key", transcription, "English")
    worker.api_keys = ["test-key"]
    monkeypatch.setattr(translation_processor.openai, "OpenAI", lambda **_kwargs: object())
    monkeypatch.setattr(worker, "_call_openai", lambda *_args, **_kwargs: translated)
    worker.finished.connect(lambda ok, text: captured.append((ok, text)))
    worker.run()
    assert captured == [(True, translated)]

    export_paths = iter(
        [
            (str(tmp_path / "traduzione.txt"), ""),
            (str(tmp_path / "traduzione.docx"), ""),
        ]
    )
    monkeypatch.setattr(
        translation_dialog.QFileDialog,
        "getSaveFileName",
        lambda *_args, **_kwargs: next(export_paths),
    )
    monkeypatch.setattr(
        translation_dialog.QMessageBox,
        "information",
        lambda *_args, **_kwargs: None,
    )
    dialog = _TranslationDialogStub(translated)
    translation_dialog.TranslationDialog.esporta_risultato(dialog, "txt")
    translation_dialog.TranslationDialog.esporta_risultato(dialog, "docx")

    assert (tmp_path / "traduzione.txt").read_text(encoding="utf-8") == translated
    assert translated in "\n".join(
        p.text for p in Document(tmp_path / "traduzione.docx").paragraphs
    )


def test_semantic_genealogy_exports_gedcom_and_both_csv_files(tmp_path):
    payload = {
        "metadata": {"comunita": "Trento", "anno": "1882"},
        "atti": [
            {
                "tipo": "nascita",
                "soggetto": {
                    "nome": "Giovanni",
                    "cognome": "Rossi",
                    "sesso": "M",
                    "data_nascita": "4 marzo 1882",
                    "luogo_nascita": "Trento",
                },
                "padre": {"nome": "Luigi", "cognome": "Rossi"},
                "madre": {"nome": "Maria", "cognome_nubile": "Bianchi"},
            }
        ],
    }
    generator = GedcomGenerator(source_system="ATK-Pro_PreRC")
    assert generator.process_ai_json(payload) is True
    gedcom_path = tmp_path / "genealogia.ged"

    generator.save_to_file(str(gedcom_path))

    gedcom = gedcom_path.read_text(encoding="utf-8")
    assert "1 NAME Giovanni /Rossi/" in gedcom
    assert "1 NAME Luigi /Rossi/" in gedcom
    assert "1 NAME Maria /Bianchi/" in gedcom
    assert gedcom.rstrip().endswith("0 TRLR")
    assert (tmp_path / "genealogia_REVISIONE.csv").exists()
    assert (tmp_path / "genealogia_REGISTRO_ORIGINALE.csv").exists()


def test_ocr_review_cancellation_does_not_publish_outputs(tmp_path, monkeypatch):
    source = tmp_path / "annullato.png"
    Image.new("RGB", (32, 32), "white").save(source)
    worker = AdvancedOCRWorker(
        "OpenAI",
        "test-key",
        ["txt", "docx", "xml"],
        str(tmp_path),
    )
    worker.api_keys = ["test-key"]
    monkeypatch.setattr(worker, "_transcribe_image", lambda *_: "testo provvisorio")

    worker.process_file(str(source), review_callback=lambda *_: None)

    assert list(tmp_path.glob("annullato_trascrizione.*")) == []


def test_empty_translation_is_not_exported(tmp_path, monkeypatch):
    warnings = []
    monkeypatch.setattr(
        translation_dialog.QMessageBox,
        "warning",
        lambda *_args: warnings.append(True),
    )
    monkeypatch.setattr(
        translation_dialog.QFileDialog,
        "getSaveFileName",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("Il selettore file non deve aprirsi")
        ),
    )

    translation_dialog.TranslationDialog.esporta_risultato(
        _TranslationDialogStub("   "),
        "txt",
    )

    assert warnings == [True]
    assert list(tmp_path.iterdir()) == []
